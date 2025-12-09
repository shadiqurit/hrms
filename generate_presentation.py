from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()

    # defined layouts
    title_slide_layout = prs.slide_layouts[0]
    bullet_slide_layout = prs.slide_layouts[1]

    # Slide Data
    slides_content = [
        {
            "type": "title",
            "title": "HR / Employee Related MIS Reports",
            "subtitle": "Prepared for: Management & HR Department\nPrepared by: [Your Name]"
        },
        {
            "type": "bullet",
            "title": "Introduction",
            "content": [
                "MIS = Management Information System",
                "HR MIS provides accurate, real-time insights about employees.",
                "Helps in decision-making, planning, workforce management.",
                "Ensures efficiency, transparency and compliance in HR processes."
            ]
        },
        {
            "type": "bullet",
            "title": "Importance of HR MIS",
            "content": [
                "Centralized employee database",
                "Reduces manual work & errors",
                "Improves HR productivity",
                "Supports data-driven decisions",
                "Ensures organizational transparency",
                "Provides analytical trends & HR KPIs"
            ]
        },
        {
            "type": "bullet",
            "title": "Key Categories of HR MIS Reports",
            "content": [
                "Employee Information Reports",
                "Attendance & Leave Reports",
                "Payroll & Compensation Reports",
                "Performance Management Reports",
                "Recruitment Reports",
                "Training & Development Reports",
                "HR Compliance Reports",
                "Employee Benefits Reports"
            ]
        },
        {
            "type": "bullet",
            "title": "Employee Information Reports",
            "content": [
                "Employee Master List",
                "Department-wise Employee List",
                "New Joiner Report",
                "Transfer & Promotion History",
                "Resignation / Separation Report",
                "Employee Contact & Profile Summary"
            ]
        },
        {
            "type": "bullet",
            "title": "Attendance & Leave MIS",
            "content": [
                "Daily Attendance Summary",
                "Monthly Presence/Absence Report",
                "Late/ Early Leave Reports",
                "Leave Balance Summary",
                "Overtime (OT) Report",
                "Shift-wise Attendance",
                "Leave Trend Analysis"
            ]
        },
        {
            "type": "bullet",
            "title": "Payroll MIS Reports",
            "content": [
                "Monthly Salary Summary",
                "Allowances & Deductions",
                "Overtime Payment",
                "Department-wise Salary Cost",
                "Tax, PF & Other Statutory Deductions",
                "Payroll Variance Report",
                "Salary Arrears & Adjustments"
            ]
        },
        {
            "type": "bullet",
            "title": "Performance & Appraisal Reports",
            "content": [
                "KPI Score Report",
                "Employee Appraisal Summary",
                "High & Low Performer Dashboard",
                "Probation Assessment Report",
                "Departmental Performance Comparison"
            ]
        },
        {
            "type": "bullet",
            "title": "Recruitment MIS Reports",
            "content": [
                "Vacancy Status Report",
                "Applicant Tracking Report",
                "Interview Evaluation Summary",
                "Selected vs Rejected Candidates",
                "Recruitment Lead Time Analysis"
            ]
        },
        {
            "type": "bullet",
            "title": "Training & Development Reports",
            "content": [
                "Annual Training Calendar",
                "Training Attendance",
                "Completed Training List",
                "Employee Skill Matrix",
                "Training Effectiveness Report"
            ]
        },
        {
            "type": "bullet",
            "title": "HR Compliance & Discipline Reports",
            "content": [
                "Warning / Show Cause Reports",
                "Disciplinary Action History",
                "Grievance Handling Summary",
                "Policy Compliance Status"
            ]
        },
        {
            "type": "bullet",
            "title": "Employee Benefits Reports",
            "content": [
                "PF / Gratuity Eligibility List",
                "Medical Benefit Usage",
                "Insurance Enrollment Report",
                "HR Loan Summary"
            ]
        },
        {
            "type": "bullet",
            "title": "HR MIS Dashboard (APEX/ERP)",
            "content": [
                "Key KPIs to Display:",
                "Total Active Employees",
                "New Joiners vs Resignations",
                "Department-wise Headcount",
                "Attendance Overview (P/L/A)",
                "Monthly Salary Cost Trend",
                "Leave Utilization Chart"
            ]
        },
        {
            "type": "bullet",
            "title": "Benefits of Implementing HR MIS Dashboard",
            "content": [
                "Better decision-making",
                "Higher transparency",
                "Reduced HR workload",
                "Improved employee satisfaction",
                "Real-time insights",
                "Enhanced compliance & control"
            ]
        },
        {
            "type": "bullet",
            "title": "Conclusion",
            "content": [
                "HR MIS reports help organizations to:",
                "Monitor workforce effectively",
                "Improve productivity",
                "Optimize HR operations",
                "Support strategic planning",
                "A well-designed HR MIS is essential for modern HR management."
            ]
        }
    ]

    # Generate Slides
    for slide_data in slides_content:
        if slide_data["type"] == "title":
            slide = prs.slides.add_slide(title_slide_layout)
            title = slide.shapes.title
            subtitle = slide.placeholders[1]
            title.text = slide_data["title"]
            subtitle.text = slide_data["subtitle"]
        
        elif slide_data["type"] == "bullet":
            slide = prs.slides.add_slide(bullet_slide_layout)
            title = slide.shapes.title
            title.text = slide_data["title"]
            
            # Add body text
            tf = slide.shapes.placeholders[1].text_frame
            tf.text = slide_data["content"][0]  # First bullet
            
            for item in slide_data["content"][1:]:
                p = tf.add_paragraph()
                p.text = item
                p.level = 0

    prs.save('HR_MIS_Reports.pptx')
    print("Presentation saved as 'HR_MIS_Reports.pptx'")

if __name__ == "__main__":
    create_presentation()